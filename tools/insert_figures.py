# -*- coding: utf-8 -*-
"""
insert_figures.py — 계획서 docx 에 도면 3장을 삽입한다.

계획서_수정사항.md ⑤항 작업. 삽입 위치는 다음 세 곳이다.

  3.1절  [그림 1] 시스템 블록도    ← fig1_block_diagram.png
  6.8절  [그림 2] 프로그램 흐름도  ← fig2_program_flow.png
  6.1절  회로 구성 (그림 없음)     ← fig3_schematic.png  (신규)

"삽입 예정" 지시문은 제출본에 남으면 미완성으로 보이므로 함께 지운다.

Word 는 SVG 삽입이 판본에 따라 불안정하므로 PNG 를 쓴다.
PNG 가 없으면 sh tools/render_figures.sh 를 먼저 실행할 것.

실행:  python tools/insert_figures.py
입력:  LED_PEDD_계획서_최종본.docx
출력:  같은 파일을 갱신 (실행 전 backup/ 에 타임스탬프 사본을 남긴다)
"""

import os
import shutil
import datetime

import docx
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOC = "LED_PEDD_계획서_최종본.docx"
PNG = os.path.join("figures", "png")

# (문단을 찾는 조건, 그림 파일, 폭 cm, 캡션)
# 폭은 A4 기본 여백 기준 본문 폭(약 16cm)을 넘지 않게 잡는다.
TARGETS = [
    {"match": "[그림 1]", "png": "fig1_block_diagram.png",
     "width": 15.5, "caption": "[그림 1] 시스템 블록도"},
    {"match": "[그림 2]", "png": "fig2_program_flow.png",
     "width": 11.5, "caption": "[그림 2] 프로그램 흐름도"},
]


def clear_runs(p):
    for r in list(p.runs):
        r._element.getparent().remove(r._element)


def add_paragraph_after(p, text, italic=False, size_pt=9):
    """p 바로 뒤에 새 문단을 만든다. python-docx 에 직접 API 가 없다."""
    from docx.text.paragraph import Paragraph
    new_el = p._p.makeelement(p._p.tag, {})
    p._p.addnext(new_el)
    np = Paragraph(new_el, p._parent)
    np.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = np.add_run(text)
    run.italic = italic
    run.font.size = docx.shared.Pt(size_pt)
    return np


def put_image(p, png_path, width_cm, caption):
    """문단 내용을 그림으로 교체하고 아래에 캡션을 붙인다."""
    clear_runs(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(png_path, width=Cm(width_cm))
    add_paragraph_after(p, caption)


def main():
    if not os.path.isdir(PNG):
        raise SystemExit("PNG 폴더가 없다. sh tools/render_figures.sh 를 먼저 실행할 것.")

    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    if not os.path.isdir("backup"):
        os.makedirs("backup")
    bak = os.path.join("backup", "LED_PEDD_계획서_그림삽입전_%s.docx" % stamp)
    shutil.copy2(DOC, bak)
    print("백업: %s" % bak)

    d = docx.Document(DOC)
    ps = d.paragraphs
    done = 0

    # --- 그림 1, 2 : "삽입 예정" 자리표시 문단을 교체 ---
    for t in TARGETS:
        hit = None
        for p in ps:
            if t["match"] in p.text and "삽입" in p.text:
                hit = p
                break
        if hit is None:
            print("건너뜀: %s 자리표시 문단을 찾지 못했다 (이미 삽입됨?)" % t["match"])
            continue
        path = os.path.join(PNG, t["png"])
        if not os.path.exists(path):
            print("건너뜀: %s 없음" % path)
            continue
        put_image(hit, path, t["width"], t["caption"])
        print("삽입: %s ← %s" % (t["caption"], t["png"]))
        done += 1

    # --- 그림 3 : 6.1절 본문 끝에 회로도를 새로 추가 ---
    path3 = os.path.join(PNG, "fig3_schematic.png")
    if os.path.exists(path3):
        start = None
        for i, p in enumerate(d.paragraphs):
            if p.text.strip().startswith("6.1") and p.style.name.startswith("Heading"):
                start = i
                break
        if start is None:
            print("건너뜀: 6.1절을 찾지 못했다")
        else:
            # 6.2절 직전 문단 뒤에 넣는다
            end = start + 1
            allp = d.paragraphs
            while end < len(allp) and not (
                    allp[end].text.strip().startswith("6.2")
                    and allp[end].style.name.startswith("Heading")):
                end += 1
            anchor = allp[end - 1]
            if "그림 3" in "".join(x.text for x in allp[start:end]):
                print("건너뜀: 6.1절에 이미 그림 3이 있다")
            else:
                from docx.text.paragraph import Paragraph
                new_el = anchor._p.makeelement(anchor._p.tag, {})
                anchor._p.addnext(new_el)
                np = Paragraph(new_el, anchor._parent)
                np.alignment = WD_ALIGN_PARAGRAPH.CENTER
                np.add_run().add_picture(path3, width=Cm(15.5))
                add_paragraph_after(np, "[그림 3] 전체 시스템 회로도")
                print("삽입: [그림 3] 전체 시스템 회로도 ← fig3_schematic.png")
                done += 1

    d.save(DOC)
    print("저장 완료: %s (그림 %d건)" % (DOC, done))


if __name__ == "__main__":
    main()
