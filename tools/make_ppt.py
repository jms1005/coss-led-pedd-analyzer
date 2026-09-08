# -*- coding: utf-8 -*-
"""
make_ppt.py — 예선 보고서 PPT 초안 생성기

제출물 「가. 보고서(PPT)」의 21장 골격을 실제 pptx 파일로 만든다.
구성은 제출물/예선_PPT_골격.md 를 따르며, 운영설명 PDF 「4-가」의 7개
필수 항목과 1:1 대응한다.

실험 결과가 필요한 장(슬라이드 11·12)은 붉은 안내 상자를 넣어 비워둔다.
측정이 끝난 뒤 그 상자만 그래프로 교체하면 된다.

실행:  python tools/make_ppt.py
출력:  제출물/가_예선보고서_초안.pptx

주의: 도면 PNG 가 먼저 있어야 한다. 없으면 sh tools/render_figures.sh 실행.
"""

import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

# ---------------------------- 공통 설정 ----------------------------

FONT = "맑은 고딕"
INK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
ACCENT = RGBColor(0x1D, 0x4E, 0xD8)
WARN = RGBColor(0xB9, 0x1C, 0x1C)
BAND = RGBColor(0xF1, 0x5C, 0x22)

SW, SH = 13.333, 7.5  # 16:9 인치


def set_font(run, size, bold=False, color=INK):
    """한글이 깨지지 않도록 라틴·동아시아 글꼴을 모두 지정한다.

    run.font.name 은 <a:latin> 만 건드린다. 한글은 <a:ea> 를 따르므로
    이것을 함께 지정하지 않으면 PowerPoint 가 기본 글꼴로 대체해버린다.
    """
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT  # <a:latin>

    rPr = run.font._element  # 이미 rPr 요소다
    for tag in ("ea", "cs"):
        q = qn("a:" + tag)
        el = rPr.find(q)
        if el is None:
            el = rPr.makeelement(q, {})
            rPr.append(el)
        el.set("typeface", FONT)


def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def para(tf, text, size=16, bold=False, color=INK, space_after=6,
         level=0, first=False, align=PP_ALIGN.LEFT):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.level = level
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    set_font(run, size, bold, color)
    return p


def rect(slide, x, y, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def new_slide(prs, title, tag=None, num=None):
    """제목 띠가 있는 기본 슬라이드."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, 1.02, RGBColor(0xF8, 0xFA, 0xFC))
    rect(s, 0, 1.00, SW, 0.045, BAND)

    tf = textbox(s, 0.55, 0.16, 9.6, 0.7)
    para(tf, title, 26, True, INK, 0, first=True)

    if tag:
        tf2 = textbox(s, 10.2, 0.26, 2.7, 0.5)
        para(tf2, tag, 13, True, BAND, 0, first=True, align=PP_ALIGN.RIGHT)
    if num is not None:
        tf3 = textbox(s, 12.4, 6.92, 0.7, 0.4)
        para(tf3, str(num), 11, False, MUTED, 0, first=True, align=PP_ALIGN.RIGHT)
    return s


def bullets(slide, items, x=0.75, y=1.45, w=11.9, size=17):
    """items: (텍스트, 레벨, 강조여부) 튜플 목록."""
    tf = textbox(slide, x, y, w, SH - y - 0.6)
    for i, (text, lvl, strong) in enumerate(items):
        color = ACCENT if strong else INK
        para(tf, text, size - lvl * 2, strong, color,
             space_after=9 if lvl == 0 else 5, level=lvl, first=(i == 0))
    return tf


def note_box(slide, lines, x=0.75, y=5.55, w=11.9, h=1.3,
             fill=RGBColor(0xFF, 0xFB, 0xEB), line=RGBColor(0xA1, 0x62, 0x07),
             color=RGBColor(0x71, 0x3F, 0x12)):
    rect(slide, x, y, w, h, fill, line)
    tf = textbox(slide, x + 0.22, y + 0.13, w - 0.44, h - 0.26)
    for i, t in enumerate(lines):
        para(tf, t, 13, i == 0, color, 3, first=(i == 0))


def todo_box(slide, lines, y=2.1, h=3.2):
    """실험 후 채울 자리."""
    note_box(slide, lines, y=y, h=h,
             fill=RGBColor(0xFE, 0xF2, 0xF2),
             line=WARN, color=WARN)


def table(slide, headers, rows, x=0.75, y=1.5, w=11.9, h=None, fs=13):
    nr, nc = len(rows) + 1, len(headers)
    h = h or min(0.42 * nr + 0.1, SH - y - 0.5)
    shape = slide.shapes.add_table(nr, nc, Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    tbl = shape.table
    for c, head in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.text = ""
        para(cell.text_frame, head, fs, True, RGBColor(0xFF, 0xFF, 0xFF), 0, first=True)
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(0x33, 0x41, 0x55)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = ""
            strong = val.startswith("*")
            para(cell.text_frame, val.lstrip("*"), fs, strong,
                 ACCENT if strong else INK, 0, first=True)
            cell.fill.solid()
            cell.fill.fore_color.rgb = (RGBColor(0xFF, 0xFF, 0xFF) if r % 2
                                        else RGBColor(0xF1, 0xF5, 0xF9))
    return tbl


def picture(slide, path, x, y, w=None, h=None):
    if not os.path.exists(path):
        tf = textbox(slide, x, y, 6, 0.6)
        para(tf, "[그림 없음: %s — sh tools/render_figures.sh 실행]" % path,
             13, True, WARN, 0, first=True)
        return None
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


# ---------------------------- 슬라이드 조립 ----------------------------

PNG = os.path.join("figures", "png")


def build():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    # ---- 표지 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, SW, SH, RGBColor(0x0F, 0x17, 0x2A))
    rect(s, 0, 3.42, SW, 0.06, BAND)
    tf = textbox(s, 1.1, 1.75, 11.2, 1.6)
    para(tf, "LED 다파장 재이용 광검출 기반", 34, True, RGBColor(0xFF, 0xFF, 0xFF), 4, first=True)
    para(tf, "저비용 액체 광학 분석기", 34, True, RGBColor(0xFF, 0xFF, 0xFF), 0)
    tf = textbox(s, 1.1, 3.75, 11.2, 0.6)
    para(tf, "포토다이오드 없이, LED만으로 액체를 판별한다", 18, False,
         RGBColor(0xF1, 0x5C, 0x22), 0, first=True)
    tf = textbox(s, 1.1, 5.1, 11.2, 1.6)
    para(tf, "2026 COSS 차세대반도체 MCU 응용 경진대회", 15, False,
         RGBColor(0x94, 0xA3, 0xB8), 6, first=True)
    para(tf, "팀명: (기입)   |   팀원: (4인 성명 기입)   |   소속: (기입)", 15,
         False, RGBColor(0x94, 0xA3, 0xB8), 6)
    para(tf, "MCU: ATmega328P   |   개발환경: Microchip Studio (레지스터 직접 제어)",
         15, False, RGBColor(0x94, 0xA3, 0xB8), 0)

    n = [1]

    def num():
        n[0] += 1
        return n[0]

    # ---- 1. 개요 ----
    s = new_slide(prs, "문제 제기 — 왜 LED로 검출하는가", "1. 개요", num())
    bullets(s, [
        ("기존 액체 광학 분석의 구성", 0, True),
        ("포토다이오드 + TIA(트랜스임피던스 증폭) + 정밀 OP-AMP + ADC", 1, False),
        ("부품 수가 늘고, 아날로그 프론트엔드 설계 난이도가 높다", 1, False),
        ("LED는 이미 광검출 소자다", 0, True),
        ("LED는 p-n 접합 소자다. 역방향 바이어스 상태에서는 포토다이오드와", 1, False),
        ("동일한 원리로 입사광에 비례하는 광전류를 만든다", 1, False),
        ("→ 전용 검출 소자 없이 LED 자체로 검출이 가능하다", 1, True),
        ("본 설계의 선택", 0, True),
        ("아날로그 증폭·ADC를 쓰지 않고, 방전 시간이라는 시간 영역 값으로 변환한다", 1, False),
        ("MCU의 디지털 핀과 내장 주변장치만으로 측정이 완결된다", 1, False),
    ])
    note_box(s, ["계획서 1.1절. 경진대회 평가 항목 중 경제성과 직결되는 선택이다."])

    s = new_slide(prs, "한 장 요약 — 최종 데모 시나리오", "1. 개요", num())
    bullets(s, [
        ("① 시료가 담긴 큐벳을 차광 챔버에 삽입한다", 0, False),
        ("② 측정 시작 버튼을 누른다", 0, False),
        ("③ MCU가 Red → Green → Blue 순으로 발광 LED를 점등한다", 0, False),
        ("④ 각 파장에 대해 검출 LED의 충전–방전 시간을 Timer1으로 측정한다", 0, False),
        ("⑤ 3채널 측정값으로 광학 지문 벡터 [R, G, B]를 구성한다", 0, False),
        ("⑥ MCU 내부에서 암전류 보정·정규화 후 최근접 중심점 분류를 실행한다", 0, False),
        ("⑦ 판정 결과를 OLED에 표시한다 (UART로 동시 로깅)", 0, True),
    ], y=1.55, size=18)
    note_box(s, ["★ 완성 사진 또는 시연 영상 썸네일을 이 자리에 넣을 것",
                 "심사위원이 가장 먼저 보는 장이다. '무엇을 만들었는가'가 즉시 전달되어야 한다."],
             y=5.5, h=1.3)

    # ---- 2. 기능·원리·구조 ----
    s = new_slide(prs, "PEDD 동작 원리", "2. 기능·원리·구조 〔4-가-1〕", num())
    bullets(s, [
        ("1단계 — 충전", 0, True),
        ("검출 LED의 캐소드를 HIGH로 구동해 역방향 바이어스를 건다.", 1, False),
        ("접합 정전용량(수십 pF)이 전원 전압까지 충전된다.", 1, False),
        ("2단계 — 방전", 0, True),
        ("핀을 입력(Hi-Z)으로 전환하면 광전류만이 방전 경로가 된다.", 1, False),
        ("빛이 셀수록 광전류가 크고 전압이 빨리 떨어진다.", 1, False),
        ("3단계 — 시간 측정", 0, True),
        ("전압이 내부 1.1V 밴드갭 기준을 지나는 순간을 Analog Comparator가 잡고,", 1, False),
        ("그 사건이 Timer1 Input Capture로 하드웨어에서 직접 기록된다.", 1, False),
        ("방전 시간 ∝ 1 / 광량   →  시간이라는 디지털 값으로 광량을 얻는다", 0, True),
    ], size=16)
    note_box(s, ["계획서 3.2절 / 3.6절. 논리 문턱값 대신 밴드갭 기준을 쓰는 이유는",
                 "논리 문턱값이 전원 전압과 온도에 따라 흔들려 측정 재현성을 해치기 때문이다."])

    s = new_slide(prs, "시스템 블록도", "2. 기능·원리·구조 〔4-가-1〕", num())
    # 그림 비율 1960:1040 → 폭 9.6in 이면 높이 5.09in. 안내 상자 자리를 남긴다.
    picture(s, os.path.join(PNG, "fig1_block_diagram.png"), 1.85, 1.3, w=9.6)
    note_box(s, ["핵심: 신호 처리부(Analog Comparator + Timer1 Input Capture)가 MCU 내부에 있다.",
                 "외부 아날로그 회로가 존재하지 않는다."], y=6.55, h=0.8)

    s = new_slide(prs, "다파장 광학 지문", "2. 기능·원리·구조 〔4-가-1〕", num())
    bullets(s, [
        ("단일 파장의 한계", 0, True),
        ("투과율 하나로는 '어둡다'는 사실만 알 뿐, 착색인지 탁도인지 구분하지 못한다", 1, False),
        ("3채널 벡터", 0, True),
        ("R·G·B 각각의 방전 시간을 재어 [R, G, B] 벡터를 만든다", 1, False),
        ("암전류 보정 후 합이 1000이 되도록 정규화 → 광량·거리 변동이 상쇄된다", 1, False),
        ("남는 것은 시료의 파장별 상대 응답, 즉 광학 지문이다", 1, True),
        ("선행 연구와의 차이", 0, True),
        ("기존 PEDD 연구는 대부분 1~2 파장, 단일 분석물 정량에 그쳤다 (계획서 2.2절)", 1, False),
        ("본 설계는 다파장 벡터로 시료 '종류'를 분류한다", 1, False),
    ])

    # ---- 3. 구현 전략 ----
    s = new_slide(prs, "프로그램 흐름", "3. 구현 전략 〔4-가-2〕", num())
    picture(s, os.path.join(PNG, "fig2_program_flow.png"), 4.35, 1.25, h=5.95)

    s = new_slide(prs, "MCU 선정 근거 — ATmega328P", "3. 구현 전략 〔4-가-2〕", num())
    table(s,
          ["선정 기준", "요구 사항", "ATmega328P 사양 및 적합성"],
          [["타이머", "μs 분해능 시간 측정", "*Timer1 16비트, /8에서 500ns 분해능"],
           ["비교기", "임계 통과 시점 포착", "*Analog Comparator 내장 + 1.1V 밴드갭 기준"],
           ["캡처", "지연 없는 시각 기록", "*Input Capture가 비교기 출력에 직결"],
           ["GPIO", "광학부 4 + 버튼 1 + I2C 2 + UART 1", "23개 I/O — 여유 있으나 과잉 아님"],
           ["개발환경", "레지스터 직접 제어", "Microchip Studio + avr-gcc"],
           ["가격", "저비용 구성", "IC 단품 약 5,000원"]],
          y=1.5, fs=14)
    note_box(s, ["STM32를 쓰지 않은 이유: 본 설계는 ADC를 쓰지 않으므로 고성능 ADC의 이점이 없고,",
                 "필요한 것은 '비교기 출력을 타이머가 직접 캡처하는 경로'인데 이 조합이 328P에 그대로 있다."],
             y=5.75, h=1.1)

    s = new_slide(prs, "부품 선정 근거", "3. 구현 전략 〔4-가-2〕", num())
    bullets(s, [
        ("발광 LED를 개별 소자 3개로 구성한 이유", 0, True),
        ("RGB 일체형은 공통 애노드/캐소드 구조라 개별 역바이어스 제어에 제약이 생긴다", 1, False),
        ("개별 소자는 파장별 광량을 저항으로 독립 조정할 수 있다", 1, False),
        ("검출 LED를 '투명 렌즈 + 적색'으로 지정한 이유", 0, True),
        ("LED는 자기 발광 파장 이하의 빛만 흡수한다 (밴드갭 이상의 광자만 흡수)", 1, False),
        ("→ 적색 검출 소자는 R·G·B를 모두 받지만, 청색 소자는 Red를 받지 못한다", 1, True),
        ("착색 렌즈는 자체가 광학 필터로 작용해 측정을 왜곡한다", 1, False),
        ("전류 제한 저항을 4종(220/330/470/1k) 확보한 이유", 0, True),
        ("광량은 Phase 1에서 반드시 조정하게 되는 변수다. 한 종류만으로는 대응이 막힌다", 1, False),
    ], size=16)

    s = new_slide(prs, "단계별 실험 계획", "3. 구현 전략 〔4-가-2〕", num())
    table(s,
          ["단계", "목표", "성공 기준", "상태"],
          [["Phase 1", "단일 채널 개념 증명", "반복 측정 CV ≤ 10%, 시료 2종 이상 구분", "★ 진행 예정"],
           ["Phase 2", "RGB 다파장 확장", "채널 추가 시 시료군 분리도 개선", "★ 진행 예정"],
           ["Phase 3", "보정 파이프라인", "암전류 보정·정규화의 유효성 확인", "★ 진행 예정"],
           ["Phase 4", "온칩 분류·시스템 통합", "PC 없이 버튼→판정→표시 전 과정 동작", "*펌웨어 구현 완료"]],
          y=1.6, fs=14, h=2.4)
    note_box(s, ["예선 시점 기준 진행 상황을 정직하게 기재할 것.",
                 "요강 1-나: '동작하는 결과 못지않게 중요한 것이 문제 해결 과정'"],
             y=4.4, h=1.0)

    # ---- 4. 설계 및 실험 과정 ----
    s = new_slide(prs, "Phase 1 결과 — 단일 채널 재현성", "4. 설계·실험 과정 〔4-가-3〕", num())
    todo_box(s, [
        "★ 실험 후 채울 것",
        "",
        "· 시료별(증류수 / 착색수 / 희석우유) 방전 시간 분포 그래프",
        "· 반복 측정 변동계수(CV) 표 — 계획서 4.1.4절 통과 기준 10% 이내",
        "· tools/analyze.py 출력 화면 캡처",
        "",
        "실행:  python tools/analyze.py data/distilled.csv data/dye.csv data/milk.csv",
        "       (스크립트는 합성 데이터로 동작 검증 완료)",
    ], y=1.55, h=4.4)

    s = new_slide(prs, "Phase 2 결과 — 다파장 분리도", "4. 설계·실험 과정 〔4-가-3〕", num())
    todo_box(s, [
        "★ 실험 후 채울 것",
        "",
        "· R·G·B 채널별 측정값과 정규화 지문 벡터",
        "· 채널 구성(R → RG → RGB)에 따른 시료군 분리도 개선 그래프",
        "· 계획서 4.2.4절 성공 기준: 채널을 늘렸을 때 가장 가까운 시료 쌍의 거리가 멀어지는가",
        "",
        "analyze.py 3번 항목이 이 값을 자동 계산한다.",
    ], y=1.55, h=4.4)

    s = new_slide(prs, "시행착오와 해결", "4. 설계·실험 과정 〔4-가-3〕", num())
    table(s,
          ["문제", "원인", "해결"],
          [["프로젝트가 열리지 않음", "componentinfo.xml 누락으로 디바이스 팩 경로가 비었음",
            "*파일 추가 후 정상 로딩"],
           ["SRAM 23.6% 점유", "AVR이 문자열 리터럴을 부팅 시 SRAM으로 복사",
            "*PROGMEM 이전 → 4.4%"],
           ["정규화 값 붕괴", "합이 작을 때 정수 나눗셈이 잘려 지문이 깨짐",
            "*곱셈 전 우측 시프트로 변경"],
           ["로그 파일 헤더 혼입", "측정마다 CSV 헤더를 재출력", "*부팅 시 1회만 출력"]],
          y=1.5, fs=13, h=2.9)
    note_box(s, ["★ 실험 중 발생한 하드웨어 문제를 여기에 추가할 것 (계획서 4.5절 실패 대응표 참조)",
                 "요강 1-나가 명시한 항목이다. 문제를 겪지 않은 척하는 것보다 해결 과정을 보이는 편이 유리하다."],
             y=4.65, h=1.1)

    s = new_slide(prs, "예외 상황 대비", "4. 설계·실험 과정 〔4-가-3〕", num())
    bullets(s, [
        ("① 16비트 타이머의 32.8ms 한계 → Overflow Counter", 0, True),
        ("Timer1은 65,536틱에서 넘친다. 오버플로우 인터럽트로 횟수를 세어", 1, False),
        ("32비트로 확장했다. 측정 상한 61회 오버플로우 ≈ 2.0초", 1, False),
        ("② 빛이 오지 않는 경우 → TIMEOUT 반환", 0, True),
        ("상한을 넘으면 측정을 중단하고 상태를 TIMEOUT으로 돌려준다", 1, False),
        ("화면에 어느 채널이 실패했는지 표시한다 (TIMEOUT: R)", 1, False),
        ("→ 채널별로 원인이 갈린다. 특정 채널만 실패하면 검출 LED 색상 문제", 1, True),
        ("③ 학습 데이터가 없는 상태 → NOT TRAINED 안내", 0, True),
        ("EEPROM의 매직·버전·체크섬을 모두 검사한다. 공장 출하 상태(0xFF)는 자연히 걸러진다", 1, False),
        ("④ OLED 미응답 → UART 로깅은 계속", 0, True),
        ("I2C 대기 루프에 상한을 두어, 표시장치 고장이 측정을 멈추지 않게 했다", 1, False),
    ], size=15)

    # ---- 5. 단위별 상세 ----
    s = new_slide(prs, "회로도", "5. 단위별 상세 〔4-가-4〕", num())
    # 그림 비율 2960:2500 → 높이 5.7in 이면 폭 6.75in. 오른쪽에 설명 패널을 둔다.
    picture(s, os.path.join(PNG, "fig3_schematic.png"), 0.5, 1.3, h=5.7)
    note_box(s, ["검출 LED(D4)만 극성이 반대다",
                 "캐소드가 PD7(AIN1)에 연결된다. 반대로 꽂으면 순방향이 되어",
                 "전하가 축적되지 않고 측정이 아예 성립하지 않는다.",
                 "",
                 "TP1(PB5) — 사이클 시간 측정용 타이밍 출력",
                 "측정 구간에만 HIGH가 된다. 보드 내장 LED와 공유하므로",
                 "부품이 늘지 않는다.",
                 "",
                 "SW1은 내부 풀업을 써서 외부 저항이 없다.",
                 "OLED 모듈은 SDA/SCL 풀업을 내장한다."],
             x=7.55, y=1.35, w=5.3, h=3.6)
    note_box(s, ["원본 도면은 제출물/나_회로도.pdf (A3 가로 1페이지)로 별도 제출한다."],
             x=7.55, y=5.2, w=5.3, h=0.7,
             fill=RGBColor(0xF1, 0xF5, 0xF9), line=RGBColor(0x94, 0xA3, 0xB8),
             color=RGBColor(0x33, 0x41, 0x55))

    s = new_slide(prs, "레지스터 직접 제어", "5. 단위별 상세 〔4-가-4〕", num())
    bullets(s, [
        ("Arduino IDE와 외부 라이브러리를 전혀 사용하지 않았다", 0, True),
        ("개발환경은 Microchip Studio, 모든 주변장치를 레지스터로 직접 다룬다", 1, False),
    ], y=1.4, size=16)
    table(s,
          ["주변장치", "직접 제어한 레지스터", "용도"],
          [["Analog Comparator", "ACSR, ADCSRB, DIDR1", "밴드갭 기준 선택, 캡처 경로 연결"],
           ["Timer1", "TCCR1A/B, TIMSK1, ICR1, TCNT1", "500ns 분해능 입력 캡처"],
           ["GPIO", "DDRB/D, PORTB/D, PIND", "충전·방전 전환, 발광 제어, 버튼"],
           ["USART", "UBRR0, UCSR0B/C, UDR0", "38400 8N1 로깅"],
           ["TWI (I2C)", "TWBR, TWCR, TWDR, TWSR", "400kHz OLED 통신"],
           ["EEPROM", "EECR, EEDR, EEARH/L", "중심점 저장 (라이브러리 미사용)"]],
          y=2.35, fs=13, h=2.9)
    note_box(s, ["avr-libc 표준 헤더(<avr/io.h> 등)는 컴파일러 제공 헤더이며 요강이 금지한 외부 라이브러리가 아니다.",
                 "EEPROM 접근은 avr-libc 함수 대신 레지스터를 직접 다뤄 요강 취지에 맞췄다."],
             y=5.5, h=1.1)

    s = new_slide(prs, "하드웨어 캡처의 이점", "5. 단위별 상세 〔4-가-4〕", num())
    bullets(s, [
        ("소프트웨어 폴링 방식", 0, True),
        ("루프에서 핀 상태를 읽어 변화를 감지 → 인터럽트·명령어 지연이 측정값에 섞인다", 1, False),
        ("선행 연구[4]가 보고한 '펌웨어 타이밍 의존성' 문제가 여기서 발생한다", 1, False),
        ("본 설계 — 하드웨어 입력 캡처", 0, True),
        ("비교기 출력이 Timer1의 캡처 입력에 직결되어 있다", 1, False),
        ("임계 통과 순간의 타이머 값이 하드웨어에 의해 ICR1로 자동 복사된다", 1, False),
        ("CPU가 인터럽트를 늦게 처리해도 캡처된 시각 자체는 흔들리지 않는다", 1, True),
        ("하드웨어와 소프트웨어의 배분", 0, True),
        ("시간에 민감한 부분은 하드웨어(비교기+캡처)에, 판단 로직은 소프트웨어에 두었다", 1, False),
    ], size=16)
    note_box(s, ["요강 3-다 '하드웨어와 소프트웨어의 배분' 대응 항목."])

    s = new_slide(prs, "온칩 분류 구현", "5. 단위별 상세 〔4-가-4〕", num())
    bullets(s, [
        ("전 구간 32비트 정수 연산 — 부동소수점 라이브러리를 링크하지 않는다", 0, True),
        ("float를 쓰면 소프트웨어 부동소수점 루틴이 1.5~3KB 딸려 들어온다", 1, False),
        ("빌드 결과 확인: 부동소수점 심볼 0건", 1, True),
        ("파이프라인", 0, True),
        ("① 컨덕턴스 G = K / t  (K = 10⁹, 빛이 셀수록 t가 작고 G가 크다)", 1, False),
        ("② 암전류 보정  I = G − G_dark", 1, False),
        ("③ 정규화  합이 1000이 되도록 — 곱셈 전 우측 시프트로 오버플로우 회피", 1, False),
        ("④ 제곱 유클리드 거리로 최근접 중심점 선택, 임계값 초과 시 UNKNOWN", 1, False),
        ("제곱근을 계산하지 않는다 — 크기 비교만 하면 되므로 제곱거리로 충분", 1, True),
        ("EEPROM 26바이트에 중심점 3×3 + 임계값 + 매직·버전·체크섬 저장", 0, True),
    ], size=15)

    # ---- 6. 해결하지 못한 부분 ----
    s = new_slide(prs, "해결하지 못한 부분", "6. 미해결 〔4-가-5〕", num())
    bullets(s, [
        ("실험을 통해서만 확정되는 항목", 0, True),
        ("분류 임계값, 최소 신호 하한, 반복 측정 횟수는 현재 임시값이다.", 1, False),
        ("같은 시료를 반복 측정했을 때의 분산을 봐야 정할 수 있어 상수로 분리해두었다.", 1, False),
        ("하드웨어 검증이 남은 부분", 0, True),
        ("I2C·OLED·EEPROM·버튼 드라이버는 컴파일 검증까지만 되어 있다.", 1, False),
        ("분류 로직은 PC 단위 테스트 24개 항목으로 검증했으나, 드라이버는 실물이 필요하다.", 1, False),
        ("설계상 남은 한계", 0, True),
        ("검출 LED의 개체차 보정은 같은 로트 소자로 완화했을 뿐 완전히 제거하지 못했다.", 1, False),
        ("온도 드리프트 보정 미구현 — 세션 시작·종료 시 증류수 재측정 절차로 대응한다.", 1, False),
        ("적외선(940nm) 채널은 구매한 소자가 가시광 차단 패키지여서 이번 범위에서 제외했다.", 1, False),
    ], size=15)
    note_box(s, ["요강이 명시적으로 요구하는 항목이다. 비워두지 말 것.",
                 "계획서 7장: '미완료 부분을 명시하는 것이 누락하는 것보다 유리하다'"])

    # ---- 7. 메모리 사용량 ----
    s = new_slide(prs, "메모리 사용량", "7. 메모리 〔4-가-6〕", num())
    table(s,
          ["프로젝트", "역할", "Flash / 32KB", "SRAM / 2KB"],
          [["01_blink", "툴체인 검증", "176 B (0.5%)", "0 B (0.0%)"],
           ["02_phase1", "단일 채널 측정", "1,430 B (4.4%)", "89 B (4.3%)"],
           ["03_phase2", "RGB 다파장 스캔", "1,360 B (4.2%)", "121 B (5.9%)"],
           ["04_phase4", "*최종 시스템", "*6,690 B (20.4%)", "*90 B (4.4%)"]],
          y=1.55, fs=14, h=2.4)
    bullets(s, [
        ("EEPROM 사용량: 26 B / 1,024 B (2.5%) — 중심점 3×3 + 임계값 + 무결성 검사", 0, False),
        ("SRAM을 4.4%로 유지한 방법: 화면 문구 전부를 Flash(PROGMEM)에 상주시켰다.", 0, False),
    ], y=4.15, size=15)
    note_box(s, ["★ Microchip Studio 빌드 출력 화면 캡처를 이 자리에 넣을 것 — 요강이 캡처 이미지를 요구한다."],
             y=5.5, h=0.9)

    # ---- 8. 부품 가격표 ----
    s = new_slide(prs, "부품 가격표와 경제성", "8. 가격 〔4-가-7〕", num())
    table(s,
          ["구분", "기존 포토다이오드 방식", "제안 방식 (PEDD)"],
          [["광원", "RGB LED 3개", "RGB LED 3개"],
           ["검출 소자", "포토다이오드 1개", "*LED 1개"],
           ["증폭 회로", "TIA용 OP-AMP 1개 + 정밀 저항", "*없음"],
           ["AD 변환", "MCU 내장 ADC 사용", "*없음 (시간 측정으로 대체)"],
           ["MCU 사용 핀", "광원 3 + ADC 1 + 기타", "*9핀 (표시·로깅 포함)"]],
          y=1.5, fs=14, h=2.6)
    todo_box(s, [
        "★ 실제 구매가로 갱신할 것 (계획서 5.1절 BOM)",
        "· 결제 내역의 실제 단가를 기입 — 추정가보다 실제 금액이 유리하다",
        "· 절감률 = (기존 − 제안) / 기존 × 100",
        "· 시제품은 Uno 보드로 구현했으나 원가는 IC 단품 구성 기준으로 산출했음을 한 줄 명시",
    ], y=4.35, h=1.75)

    # ---- 마지막: 제출 전 확인 ----
    s = new_slide(prs, "제출 전 최종 확인", "체크리스트", num())
    bullets(s, [
        ("□ 보고서 PPT", 0, False),
        ("□ 회로도 PDF — 제출물/나_회로도.pdf (A3 가로 1페이지, 생성 완료)", 0, False),
        ("□ 소스코드 — firmware/ 전체", 0, False),
        ("□ 동작 동영상", 0, False),
        ("□ 팀원 4명 전원 구글폼 개별 제출 — 누락 시 실격", 0, True),
        ("□ 메일 발송: polaris@disu.ac.kr", 0, False),
        ("마감: 2026-09-18 (금)", 0, True),
    ], y=1.7, size=18)

    out = os.path.join("제출물", "가_예선보고서_초안.pptx")
    if not os.path.isdir("제출물"):
        os.makedirs("제출물")
    prs.save(out)
    print("생성 완료: %s (슬라이드 %d장, %d bytes)"
          % (out, len(prs.slides._sldIdLst), os.path.getsize(out)))


if __name__ == "__main__":
    build()
